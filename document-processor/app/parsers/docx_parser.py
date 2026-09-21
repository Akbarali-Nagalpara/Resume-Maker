"""DOCX structural metadata extraction with python-docx.

Assists template analysis (paragraph/style inventory, tables, columns,
images, headers/footers, page geometry). Content authority for DOCX stays
with the Spring Boot POI parser; this module feeds tooling, tests and the
optional metadata endpoint.
"""

from __future__ import annotations

import hashlib
import io

from docx import Document
from docx.shared import Emu


def _twips_to_points(value) -> float:
    try:
        return float(value.twips) / 20.0
    except (AttributeError, TypeError, ValueError):
        return 0.0


def parse_docx(data: bytes, filename: str) -> dict:
    checksum = hashlib.sha256(data).hexdigest()
    document = Document(io.BytesIO(data))

    paragraphs: list[dict] = []
    for i, paragraph in enumerate(document.paragraphs):
        runs = [
            {
                "text": run.text,
                "bold": bool(run.bold),
                "italic": bool(run.italic),
                "size": run.font.size.pt if run.font.size else None,
                "color": str(run.font.color.rgb) if run.font.color else None,
                "font": run.font.name,
            }
            for run in paragraph.runs
        ]
        paragraphs.append(
            {
                "id": f"p{i}",
                "text": paragraph.text,
                "style": paragraph.style.name if paragraph.style else None,
                "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
                "runs": runs,
            }
        )

    tables = len(document.tables)
    inline_images = len(document.inline_shapes)
    headers = any(s.header.paragraphs for s in document.sections)
    footers = any(s.footer.paragraphs for s in document.sections)

    first = document.sections[0] if document.sections else None
    width = _twips_to_points(first.page_width) if first is not None else 595.0
    height = _twips_to_points(first.page_height) if first is not None else 842.0
    columns = 1
    try:
        cols = first._sectPr.xpath("./w:cols")[0] if first is not None else None
        if cols is not None and cols.get("{%s}num" % cols.nsmap.get("w", "")):
            columns = max(1, int(cols.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num")))
    except Exception:
        columns = 1

    return {
        "documentType": "DOCX",
        "sourceChecksum": checksum,
        "page": {
            "widthPoints": width or 595.0,
            "heightPoints": height or 842.0,
        },
        "paragraphCount": len(paragraphs),
        "paragraphs": paragraphs,
        "tables": tables,
        "images": inline_images,
        "columns": columns,
        "hasHeader": bool(headers),
        "hasFooter": bool(footers),
        "filename": filename,
    }
