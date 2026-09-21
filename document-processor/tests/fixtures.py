"""Representative fixtures built in-test (no binary blobs in git)."""

from __future__ import annotations

import io

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen.canvas import Canvas


def make_docx(path, *, two_column=False, multi_page=False, with_table=False) -> str:
    doc = Document()
    if two_column:
        sect = doc.sections[0]
        cols = sect._sectPr.xpath("./w:cols")[0]
        cols.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num", "2")
    name = doc.add_paragraph()
    run = name.add_run("Maya Chen")
    run.bold = True
    run.font.size = Pt(22)
    doc.add_paragraph("maya.chen@example.com | San Francisco, CA")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Product-minded engineer with 7+ years of experience.")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("TypeScript | React | Java")
    doc.add_heading("Work Experience", level=1)
    doc.add_paragraph("Senior Product Engineer @ Northstar Labs")
    doc.add_paragraph("2022 \u2014 Present")
    doc.add_paragraph("\u2022 Led delivery of a platform.")
    if with_table:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Year"
        table.cell(0, 1).text = "Role"
        table.cell(1, 0).text = "2022"
        table.cell(1, 1).text = "Engineer"
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.S. Computer Science")
    if multi_page:
        for i in range(60):
            doc.add_paragraph(f"Additional background line {i}.")
    doc.save(path)
    return path


def make_pdf(path, *, two_column=False, multi_page=False, image=False) -> str:
    c = Canvas(path, pagesize=LETTER)
    y = 720

    def line(text, font="Helvetica", size=10, x=72):
        nonlocal y
        c.setFont(font, size)
        c.drawString(x, y, text)
        y -= size + 4

    line("Maya Chen", "Helvetica-Bold", 20)
    line("maya.chen@example.com | San Francisco, CA")
    y -= 8
    line("SUMMARY", "Helvetica-Bold", 11)
    line("Product-minded engineer with 7+ years of experience.")
    y -= 8
    line("SKILLS", "Helvetica-Bold", 11)
    if two_column:
        line("TypeScript | React", size=10, x=72)
        c.setFont("Helvetica", 10)
        c.drawString(330, y + 14, "Java | Go")
    else:
        line("TypeScript | React | Java")
    y -= 8
    line("WORK EXPERIENCE", "Helvetica-Bold", 11)
    line("Senior Product Engineer @ Northstar Labs", "Helvetica-Bold", 10)
    line("2022 - Present")
    line("- Led delivery of a platform.")
    if image:
        from PIL import Image as PILImage

        img = PILImage.new("RGB", (60, 60), color=(239, 88, 72))
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp.write(buf.read())
            tmp_path = tmp.name
        c.drawImage(tmp_path, 450, 700, width=60, height=60)
    if multi_page:
        c.showPage()
        y = 720
        line("Page two content line.", size=10)
    c.save()
    return path
